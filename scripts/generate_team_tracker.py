#!/usr/bin/env python3
"""
Generate a Travnook Travels monthly Team Performance Tracker workbook
(Dashboard + Weekly Tracker sheets), in the same format Bindhya used for
August 2026: daily entry cells (yellow), auto-rolled-up weekly and monthly
totals, a mid-month checkpoint, and a per-agent + team status indicator.

Usage:
    python3 generate_team_tracker.py --year 2026 --month 9 \
        --agents SAHA,IYAD,RIHAB,SHILPA,ALSHIMAA,HARINDER \
        --target 25000 --team-lead Bindhya --output trackers/out.xlsx

Re-run this each month with the current roster/target to get a fresh
tracker instead of hand-building one.
"""
import argparse
import calendar
import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

NAVY = "FF1F3864"
WHITE = "FFFFFFFF"
SUBHEADER_BLUE = "FFD9E2F3"
TOTAL_BLUE = "FFBDD7EE"
ENTRY_YELLOW = "FFFFFF00"
INPUT_BLUE_FONT = "FF0000FF"
WARN_RED = "FFC00000"
GRAY = "FF808080"

FONT_NAME = "Arial"

STATUS_FORMULA = (
    'IF({pct}>=1,"\U0001F7E2 On Target",'
    'IF({pct}>=0.7,"\U0001F7E1 Watch","\U0001F534 Critical"))'
)


def build_weeks(year: int, month: int) -> list[list[datetime.date]]:
    """Group the month's working days (Mon-Sat, Sunday off) into week
    blocks. A new block starts every Monday; the first/last blocks are
    partial when the month doesn't start/end on a Mon/Sat."""
    _, days_in_month = calendar.monthrange(year, month)
    weeks: list[list[datetime.date]] = []
    current: list[datetime.date] = []
    for day in range(1, days_in_month + 1):
        d = datetime.date(year, month, day)
        if d.weekday() == 6:  # Sunday off
            continue
        if d.weekday() == 0 and current:  # Monday starts a new block
            weeks.append(current)
            current = []
        current.append(d)
    if current:
        weeks.append(current)
    return weeks


def week_label(idx: int, days: list[datetime.date]) -> str:
    month_abbr = days[0].strftime("%b")
    if len(days) == 1:
        return f"WK {idx}  ({days[0].day} {month_abbr})"
    return f"WK {idx}  ({days[0].day}-{days[-1].day} {month_abbr})"


def day_label(d: datetime.date) -> str:
    return f"{d.strftime('%b')} {d.day}\n{d.strftime('%a')}"


def style_header(cell, size=9, wrap=True):
    cell.font = Font(name=FONT_NAME, size=size, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=wrap)


def style_subheader(cell, wrap=True):
    cell.font = Font(name=FONT_NAME, size=9, bold=True, color=NAVY)
    cell.fill = PatternFill("solid", fgColor=SUBHEADER_BLUE)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=wrap)


def style_total(cell):
    cell.font = Font(name=FONT_NAME, size=9, bold=True)
    cell.fill = PatternFill("solid", fgColor=TOTAL_BLUE)
    cell.alignment = Alignment(horizontal="center", vertical="center")


def style_entry(cell):
    cell.font = Font(name=FONT_NAME, size=9)
    cell.fill = PatternFill("solid", fgColor=ENTRY_YELLOW)
    cell.alignment = Alignment(horizontal="center", vertical="center")


def style_body(cell, bold=False, center=True):
    cell.font = Font(name=FONT_NAME, size=9, bold=bold)
    if center:
        cell.alignment = Alignment(horizontal="center", vertical="center")


def generate(
    year: int,
    month: int,
    agents: list[str],
    monthly_target: float,
    team_lead: str,
    output: Path,
):
    weeks = build_weeks(year, month)
    working_days = sum(len(w) for w in weeks)
    n_agents = len(agents)
    month_name = datetime.date(year, month, 1).strftime("%B").upper()

    # ---- column layout for the Weekly Tracker sheet ----
    col = 4  # A=agent, B=daily target, C=monthly target
    week_cols = []  # (target_col, [day_cols], achieved_col, deficit_col, days)
    for w in weeks:
        target_col = col
        col += 1
        day_cols = list(range(col, col + len(w)))
        col += len(w)
        achieved_col = col
        col += 1
        deficit_col = col
        col += 1
        week_cols.append((target_col, day_cols, achieved_col, deficit_col, w))
    monthly_achieved_col = col
    monthly_deficit_col = col + 1
    pct_col = col + 2
    status_col = col + 3
    last_col = status_col

    # ---- mid-month checkpoint: weeks fully completed by the 15th ----
    checkpoint_idx = [i for i, w in enumerate(weeks) if w[-1].day <= 15]
    if not checkpoint_idx:
        checkpoint_idx = [0]
    checkpoint_days = sum(len(weeks[i]) for i in checkpoint_idx)
    checkpoint_end = weeks[checkpoint_idx[-1]][-1]
    checkpoint_target = round(monthly_target * checkpoint_days / working_days)

    wb = Workbook()

    # =========================================================
    # Weekly Tracker sheet
    # =========================================================
    wt = wb.active
    wt.title = "Weekly Tracker"
    wt.sheet_view.showGridLines = False
    wt.freeze_panes = "D6"

    wt.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    c = wt.cell(1, 1, f"TRAVNOOK TRAVELS — {month_name} {year}  WEEKLY & MONTHLY TRACKER")
    c.font = Font(name=FONT_NAME, size=14, bold=True, color=WHITE)
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(horizontal="left", vertical="center")
    wt.row_dimensions[1].height = 25.5

    wt.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    c = wt.cell(
        2, 1,
        f"Team Lead: {team_lead}   |   Monthly Target: AED {monthly_target:,.0f} per agent   "
        f"|   Working Days: {working_days}   |   \U0001F7E1 Yellow cells = daily data entry",
    )
    c.font = Font(name=FONT_NAME, size=9, color=NAVY)

    wt.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last_col)
    c = wt.cell(
        3, 1,
        f"⚠️  Mid-month checkpoint: be at AED {checkpoint_target:,.0f} per agent "
        f"(AED {checkpoint_target * n_agents:,.0f} team) by {checkpoint_end.strftime('%d %b')} "
        f"/ end of WK {checkpoint_idx[-1] + 1} to avoid a month-end rush.",
    )
    c.font = Font(name=FONT_NAME, size=9, bold=True, color=WARN_RED)

    # header rows 4-5
    wt.merge_cells(start_row=4, start_column=1, end_row=5, end_column=1)
    style_header(wt.cell(4, 1, "AGENT"))
    wt.merge_cells(start_row=4, start_column=2, end_row=5, end_column=2)
    style_header(wt.cell(4, 2, "DAILY\nTARGET"))
    wt.merge_cells(start_row=4, start_column=3, end_row=5, end_column=3)
    style_header(wt.cell(4, 3, "MONTHLY\nTARGET"))

    for i, (target_col, day_cols, achieved_col, deficit_col, w) in enumerate(week_cols, start=1):
        wt.merge_cells(start_row=4, start_column=target_col, end_row=4, end_column=deficit_col)
        style_header(wt.cell(4, target_col, week_label(i, w)))
        for c_ in range(target_col + 1, deficit_col):
            style_header(wt.cell(4, c_, None))
        style_subheader(wt.cell(5, target_col, "WK\nTARGET"))
        for dcol, d in zip(day_cols, w):
            style_subheader(wt.cell(5, dcol, day_label(d)))
        style_subheader(wt.cell(5, achieved_col, "WK\nACHIEVED"))
        style_subheader(wt.cell(5, deficit_col, "WK\nDEFICIT"))

    for label, cc in (
        ("MONTHLY\nACHIEVED", monthly_achieved_col),
        ("MONTHLY\nDEFICIT", monthly_deficit_col),
        ("% vs\nTARGET", pct_col),
        ("STATUS", status_col),
    ):
        wt.merge_cells(start_row=4, start_column=cc, end_row=5, end_column=cc)
        style_header(wt.cell(4, cc, label))

    wt.row_dimensions[4].height = 19.5
    wt.row_dimensions[5].height = 27.75

    # column widths
    wt.column_dimensions["A"].width = 12
    wt.column_dimensions["B"].width = 9
    wt.column_dimensions["C"].width = 10
    for target_col, day_cols, achieved_col, deficit_col, w in week_cols:
        wt.column_dimensions[get_column_letter(target_col)].width = 9
        for dcol in day_cols:
            wt.column_dimensions[get_column_letter(dcol)].width = 8.5
        wt.column_dimensions[get_column_letter(achieved_col)].width = 9
        wt.column_dimensions[get_column_letter(deficit_col)].width = 9
    for cc in (monthly_achieved_col, monthly_deficit_col, pct_col):
        wt.column_dimensions[get_column_letter(cc)].width = 11
    wt.column_dimensions[get_column_letter(status_col)].width = 13

    # agent rows
    first_agent_row = 6
    for i, agent in enumerate(agents):
        row = first_agent_row + i
        style_body(wt.cell(row, 1, agent), bold=True, center=False)
        wt.cell(row, 2, f"=ROUND(C{row}/{working_days},0)")
        style_body(wt.cell(row, 2))
        cell = wt.cell(row, 3, monthly_target)
        cell.font = Font(name=FONT_NAME, size=9, color=INPUT_BLUE_FONT)
        cell.alignment = Alignment(horizontal="center", vertical="center")

        achieved_refs = []
        for target_col, day_cols, achieved_col, deficit_col, w in week_cols:
            t_letter = get_column_letter(target_col)
            a_letter = get_column_letter(achieved_col)
            d_letter = get_column_letter(deficit_col)
            wt.cell(row, target_col, f"=ROUND($C{row}*{len(w)}/{working_days},0)")
            style_body(wt.cell(row, target_col))
            for dcol in day_cols:
                style_entry(wt.cell(row, dcol, 0))
            first_day = get_column_letter(day_cols[0])
            last_day = get_column_letter(day_cols[-1])
            wt.cell(row, achieved_col, f"=SUM({first_day}{row}:{last_day}{row})")
            style_body(wt.cell(row, achieved_col))
            wt.cell(row, deficit_col, f"={t_letter}{row}-{a_letter}{row}")
            style_body(wt.cell(row, deficit_col))
            achieved_refs.append(f"{a_letter}{row}")

        ma_letter = get_column_letter(monthly_achieved_col)
        md_letter = get_column_letter(monthly_deficit_col)
        pct_letter = get_column_letter(pct_col)
        wt.cell(row, monthly_achieved_col, "=" + "+".join(achieved_refs))
        style_body(wt.cell(row, monthly_achieved_col))
        wt.cell(row, monthly_deficit_col, f"=C{row}-{ma_letter}{row}")
        style_body(wt.cell(row, monthly_deficit_col))
        wt.cell(row, pct_col, f"=IFERROR({ma_letter}{row}/C{row},0)")
        style_body(wt.cell(row, pct_col))
        wt.cell(row, status_col, "=" + STATUS_FORMULA.format(pct=f"{pct_letter}{row}"))
        style_body(wt.cell(row, status_col))

    # team total row
    total_row = first_agent_row + n_agents
    style_total(wt.cell(total_row, 1, "TEAM TOTAL"))
    for cc in (2, 3):
        letter = get_column_letter(cc)
        wt.cell(total_row, cc, f"=SUM({letter}{first_agent_row}:{letter}{total_row - 1})")
        style_total(wt.cell(total_row, cc))
    for target_col, day_cols, achieved_col, deficit_col, w in week_cols:
        for cc in (target_col, achieved_col, deficit_col):
            letter = get_column_letter(cc)
            wt.cell(total_row, cc, f"=SUM({letter}{first_agent_row}:{letter}{total_row - 1})")
            style_total(wt.cell(total_row, cc))
    for cc in (monthly_achieved_col, monthly_deficit_col):
        letter = get_column_letter(cc)
        wt.cell(total_row, cc, f"=SUM({letter}{first_agent_row}:{letter}{total_row - 1})")
        style_total(wt.cell(total_row, cc))
    ma_letter = get_column_letter(monthly_achieved_col)
    c_total = f"C{total_row}"
    wt.cell(total_row, pct_col, f"=IFERROR({ma_letter}{total_row}/{c_total},0)")
    style_total(wt.cell(total_row, pct_col))
    wt.cell(total_row, status_col, "=" + STATUS_FORMULA.format(pct=f"{get_column_letter(pct_col)}{total_row}"))
    style_total(wt.cell(total_row, status_col))

    # mid-month checkpoint summary row
    note_row = total_row + 2
    wt.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=3)
    c = wt.cell(note_row, 1, f"MID-MONTH CHECKPOINT (through {checkpoint_end.strftime('%d %b')} / "
                              f"WK1–WK{checkpoint_idx[-1] + 1})")
    c.font = Font(name=FONT_NAME, size=9, bold=True, color=NAVY)
    achieved_terms = "+".join(
        f"{get_column_letter(week_cols[i][2])}{total_row}" for i in checkpoint_idx
    )
    target_terms = checkpoint_target * n_agents
    wt.merge_cells(start_row=note_row, start_column=4, end_row=note_row, end_column=last_col)
    c = wt.cell(
        note_row, 4,
        f'="Team actual: AED "&TEXT({achieved_terms},"#,##0")&"   |   Aim: AED {target_terms:,.0f} '
        f'team (AED {checkpoint_target:,.0f}/agent)"',
    )
    c.font = Font(name=FONT_NAME, size=9, bold=True, color=WARN_RED)

    # =========================================================
    # Dashboard sheet
    # =========================================================
    db = wb.create_sheet("Dashboard", 0)
    db.sheet_view.showGridLines = False

    db.merge_cells("A1:H1")
    c = db.cell(1, 1, f"TRAVNOOK TRAVELS — {month_name} {year} TEAM PERFORMANCE DASHBOARD")
    c.font = Font(name=FONT_NAME, size=14, bold=True, color=WHITE)
    c.fill = PatternFill("solid", fgColor=NAVY)
    db.row_dimensions[1].height = 25.5

    db.merge_cells("A2:H2")
    c = db.cell(
        2, 1,
        f"Team Lead: {team_lead}   |   Monthly Target: AED {monthly_target:,.0f} per agent   "
        f"|   Working Days: {working_days}   |   Mid-Month Checkpoint (by {checkpoint_end.strftime('%d %b')}): "
        f"AED {checkpoint_target:,.0f} per agent",
    )
    c.font = Font(name=FONT_NAME, size=9, color=NAVY)

    db.merge_cells("A4:H4")
    c = db.cell(4, 1, "▼  MONTHLY PERFORMANCE SUMMARY")
    c.font = Font(name=FONT_NAME, size=9, bold=True, color=NAVY)

    headers = [
        "TEAM TARGET", "TEAM ACHIEVED (MTD)", "TEAM DEFICIT (MTD)", "% ACHIEVED",
        "MID-MONTH TARGET (by 15th)", "MID-MONTH ACTUAL", "DAILY AIM (team)", "STATUS",
    ]
    for i, h in enumerate(headers, start=1):
        style_header(db.cell(5, i, h))
    db.row_dimensions[5].height = 31.5

    achieved_terms_dash = "+".join(
        f"'Weekly Tracker'!{get_column_letter(week_cols[i][2])}{total_row}" for i in checkpoint_idx
    )
    target_terms_dash = "+".join(
        f"'Weekly Tracker'!{get_column_letter(week_cols[i][0])}{total_row}" for i in checkpoint_idx
    )
    db.cell(6, 1, f"='Weekly Tracker'!C{total_row}")
    db.cell(6, 2, f"='Weekly Tracker'!{ma_letter}{total_row}")
    db.cell(6, 3, "=A6-B6")
    db.cell(6, 4, "=IFERROR(B6/A6,0)")
    db.cell(6, 5, "=" + target_terms_dash)
    db.cell(6, 6, "=" + achieved_terms_dash)
    db.cell(6, 7, f"=ROUND(A6/{working_days},0)")
    db.cell(6, 8, "=" + STATUS_FORMULA.format(pct="D6"))
    for i in range(1, 9):
        style_body(db.cell(6, i))

    db.merge_cells("A8:H8")
    c = db.cell(8, 1, "▼  INDIVIDUAL AGENT PERFORMANCE")
    c.font = Font(name=FONT_NAME, size=9, bold=True, color=NAVY)

    ag_headers = ["AGENT", "MONTHLY TARGET", "MTD ACHIEVED", "MTD DEFICIT", "% ACHIEVED", "DAILY AIM (agent)", "STATUS"]
    for i, h in enumerate(ag_headers, start=1):
        style_header(db.cell(9, i, h))

    for i, agent in enumerate(agents):
        row = 10 + i
        wt_row = first_agent_row + i
        style_body(db.cell(row, 1, agent), center=False)
        db.cell(row, 2, f"='Weekly Tracker'!C{wt_row}")
        db.cell(row, 3, f"='Weekly Tracker'!{ma_letter}{wt_row}")
        db.cell(row, 4, "=B{r}-C{r}".format(r=row))
        db.cell(row, 5, "=IFERROR(C{r}/B{r},0)".format(r=row))
        db.cell(row, 6, f"=ROUND(B{row}/{working_days},0)")
        db.cell(row, 7, "=" + STATUS_FORMULA.format(pct=f"E{row}"))
        for cc in range(2, 8):
            style_body(db.cell(row, cc))
        db.merge_cells(start_row=row, start_column=7, end_row=row, end_column=8)

    total_row_db = 10 + n_agents
    style_total(db.cell(total_row_db, 1, "TEAM TOTAL"))
    db.cell(total_row_db, 2, f"=SUM(B10:B{total_row_db - 1})")
    db.cell(total_row_db, 3, f"=SUM(C10:C{total_row_db - 1})")
    db.cell(total_row_db, 4, f"=B{total_row_db}-C{total_row_db}")
    db.cell(total_row_db, 5, f"=IFERROR(C{total_row_db}/B{total_row_db},0)")
    db.cell(total_row_db, 6, f"=SUM(F10:F{total_row_db - 1})")
    db.cell(total_row_db, 7, "=" + STATUS_FORMULA.format(pct=f"E{total_row_db}"))
    for cc in range(2, 8):
        style_total(db.cell(total_row_db, cc))
    db.merge_cells(start_row=total_row_db, start_column=7, end_row=total_row_db, end_column=8)

    footer_row = total_row_db + 2
    db.merge_cells(start_row=footer_row, start_column=1, end_row=footer_row, end_column=8)
    c = db.cell(
        footer_row, 1,
        f"ℹ️  Each agent's monthly target is a flat AED {monthly_target:,.0f} "
        f"({n_agents} agents × {monthly_target:,.0f} = AED {monthly_target * n_agents:,.0f} team target). "
        f"Yellow cells on 'Weekly Tracker' = daily entry; everything else here updates automatically.",
    )
    c.font = Font(name=FONT_NAME, size=9, color=GRAY)

    db.column_dimensions["A"].width = 14
    db.column_dimensions["B"].width = 16
    db.column_dimensions["C"].width = 14
    db.column_dimensions["D"].width = 14
    db.column_dimensions["E"].width = 12
    db.column_dimensions["F"].width = 16
    db.column_dimensions["G"].width = 14
    db.column_dimensions["H"].width = 14

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    print(f"Wrote {output}")


def parse_args():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--year", type=int, required=True)
    p.add_argument("--month", type=int, required=True)
    p.add_argument("--agents", type=str, required=True, help="Comma-separated agent names")
    p.add_argument("--target", type=float, required=True, help="Flat monthly target per agent (AED)")
    p.add_argument("--team-lead", type=str, required=True)
    p.add_argument("--output", type=Path, required=True)
    return p.parse_args()


if __name__ == "__main__":
    args = parse_args()
    agents = [a.strip() for a in args.agents.split(",") if a.strip()]
    generate(args.year, args.month, agents, args.target, args.team_lead, args.output)
