#!/usr/bin/env python3
"""
Generate a Travnook mid-week (default: Friday) checkpoint sheet: per-agent
weekly target, the prorated target as of a given day into the week, actual
achieved so far, deficit against that day's pace, what's still needed to
close the full week, a projected carry-forward into next week (deficits
only — surplus doesn't discount next week), a recovery-adjusted planned
target for next week, and a formula-driven remark per agent.

Usage:
    python3 generate_friday_checkpoint_report.py \
        --data "Harinder:12494,Ibrahim:11353,Layan (ATL):8600,Chaima:8041,Rihab:5874,Khosema:2550,Saha:2070,Alshima:0" \
        --default-target 25000 \
        --target-override "Layan (ATL):20000,Chaima:20000" \
        --week-label "WK2: 7-12 Sep 2026" \
        --checkpoint-day Friday --days-elapsed 5 --days-in-week 6 \
        --output trackers/reports/FRIDAY_CHECKPOINT_2026-09-11.xlsx
"""
import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

NAVY = "FF1F3864"
WHITE = "FFFFFFFF"
TOTAL_BLUE = "FFBDD7EE"
GRAY = "FF808080"
GREEN_FILL = "FFC6EFCE"
YELLOW_FILL = "FFFFEB9C"
RED_FILL = "FFFFC7CE"
FONT_NAME = "Arial"

WORKING_DAYS = 26  # per CLAUDE.md monthly working-day assumption

REMARK_FORMULA = (
    'IF(D{r}>=B{r},"\U0001F7E2 Exceeded full week target — keep pushing.",'
    'IF(D{r}>=C{r},"\U0001F7E2 On pace — close the remaining gap by Saturday.",'
    'IF(D{r}=0,"\U0001F534 No activity logged — needs a same-day check-in.",'
    'IF((C{r}-D{r})<=0.3*C{r},"\U0001F7E1 Slightly behind — recoverable by Saturday.",'
    '"\U0001F534 Significantly behind pace — needs urgent closes before Saturday."))))'
)


def style_header(cell):
    cell.font = Font(name=FONT_NAME, size=9, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_total(cell):
    cell.font = Font(name=FONT_NAME, size=9, bold=True)
    cell.fill = PatternFill("solid", fgColor=TOTAL_BLUE)
    cell.alignment = Alignment(horizontal="center", vertical="center")


def style_body(cell, bold=False, center=True):
    cell.font = Font(name=FONT_NAME, size=9, bold=bold)
    if center:
        cell.alignment = Alignment(horizontal="center", vertical="center")


def generate(
    agents: list[tuple[str, float, float]],
    week_label: str,
    checkpoint_day: str,
    days_elapsed: int,
    days_in_week: int,
    default_target: float,
    output: Path,
):
    n_agents = len(agents)
    overrides = [f"{name} AED {mt:,.0f}/mo" for name, _, mt in agents if mt != default_target]

    wb = Workbook()
    ws = wb.active
    ws.title = "Friday Checkpoint"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A10"

    last_col = 9  # A..I

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    c = ws.cell(1, 1, "TRAVNOOK TRAVEL & TOURISM — MID-WEEK CHECKPOINT")
    c.font = Font(name=FONT_NAME, size=14, bold=True, color=WHITE)
    c.fill = PatternFill("solid", fgColor=NAVY)
    ws.row_dimensions[1].height = 25.5

    override_note = f"   |   Target override: {', '.join(overrides)}" if overrides else ""
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    c = ws.cell(
        2, 1,
        f"{week_label}   |   As of: {checkpoint_day} (day {days_elapsed} of {days_in_week} in the work week)   |   "
        f"Default Monthly Target: AED {default_target:,.0f}/agent{override_note}",
    )
    c.font = Font(name=FONT_NAME, size=9, color=NAVY)

    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last_col)
    c = ws.cell(
        3, 1,
        "Carry-forward only applies to deficits (what's still short of the full week's target) — a surplus this "
        "week does not reduce next week's target.",
    )
    c.font = Font(name=FONT_NAME, size=9, italic=True, color=GRAY)

    headers = [
        "AGENT", "WEEKLY TARGET", f"TARGET BY {checkpoint_day.upper()}", f"ACHIEVED (BY {checkpoint_day.upper()})",
        f"DEFICIT vs {checkpoint_day.upper()} PACE", "STILL NEEDED (FULL WEEK)", "PROJECTED CARRY-FORWARD",
        "PLANNED TARGET — NEXT WEEK", "REMARKS",
    ]
    for i, h in enumerate(headers, start=1):
        style_header(ws.cell(5, i, h))
    ws.row_dimensions[5].height = 34

    first_row = 6
    last_row = first_row + n_agents - 1
    total_row = last_row + 1

    for i, (name, achieved, monthly_target) in enumerate(agents):
        row = first_row + i
        style_body(ws.cell(row, 1, name), center=False)
        ws.cell(row, 2, f"=ROUND({monthly_target}/{WORKING_DAYS}*{days_in_week},0)")
        ws.cell(row, 3, f"=ROUND(B{row}*{days_elapsed}/{days_in_week},0)")
        ws.cell(row, 4, achieved)
        ws.cell(row, 5, f"=C{row}-D{row}")
        ws.cell(row, 6, f"=B{row}-D{row}")
        ws.cell(row, 7, f"=MAX(F{row},0)")
        ws.cell(row, 8, f"=ROUND({monthly_target}/{WORKING_DAYS}*{days_in_week},0)+G{row}")
        ws.cell(row, 9, "=" + REMARK_FORMULA.format(r=row))
        ws.cell(row, 9).alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        for cc in (2, 3, 4, 5, 6, 7, 8):
            style_body(ws.cell(row, cc))

    style_total(ws.cell(total_row, 1, "TEAM TOTAL"))
    for cc in (2, 3, 4, 5, 6, 7, 8):
        letter = get_column_letter(cc)
        ws.cell(total_row, cc, f"=SUM({letter}{first_row}:{letter}{last_row})")
        style_total(ws.cell(total_row, cc))
    ws.cell(total_row, 9, "")
    style_total(ws.cell(total_row, 9))

    widths = [14, 13, 13, 15, 15, 14, 15, 16, 42]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in range(first_row, total_row + 1):
        ws.row_dimensions[row].height = 30

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    print(f"Wrote {output}")


def parse_pairs(s: str) -> dict[str, float]:
    result = {}
    for item in s.split(","):
        item = item.strip()
        if not item:
            continue
        name, value = item.rsplit(":", 1)
        result[name.strip()] = float(value.strip())
    return result


def parse_args():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--data", type=str, required=True, help='Comma-separated "Agent:Achieved" pairs')
    p.add_argument("--default-target", type=float, required=True)
    p.add_argument("--target-override", type=str, default="")
    p.add_argument("--week-label", type=str, required=True)
    p.add_argument("--checkpoint-day", type=str, default="Friday")
    p.add_argument("--days-elapsed", type=int, default=5)
    p.add_argument("--days-in-week", type=int, default=6)
    p.add_argument("--output", type=Path, required=True)
    return p.parse_args()


if __name__ == "__main__":
    args = parse_args()
    achieved = parse_pairs(args.data)
    overrides = parse_pairs(args.target_override) if args.target_override else {}
    triples = [
        (name, amount, overrides.get(name, args.default_target))
        for name, amount in achieved.items()
    ]
    generate(
        triples, args.week_label, args.checkpoint_day, args.days_elapsed,
        args.days_in_week, args.default_target, args.output,
    )
