#!/usr/bin/env python3
"""
Generate a Travnook Weekly Sales Report from per-agent weekly totals
(e.g. pulled from the live team dashboard). Computes team/individual
targets from the CLAUDE.md base target and working-day assumptions,
classifies each agent GREEN/YELLOW/RED, and flags who needs attention.

Usage:
    python3 generate_weekly_sales_report.py \
        --data "Harinder:12494,Ibrahim:11353,Layan (ATL):8600,Chaima:8041,Rihab:5874,Khosema:2550,Saha:2070,Alshima:0" \
        --week-ending 2026-09-12 \
        --output trackers/reports/WEEKLY_SALES_REPORT_2026-09-12.xlsx
"""
import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

NAVY = "FF1F3864"
WHITE = "FFFFFFFF"
TOTAL_BLUE = "FFBDD7EE"
INPUT_BLUE_FONT = "FF0000FF"
GRAY = "FF808080"
GREEN_FILL = "FFC6EFCE"
YELLOW_FILL = "FFFFEB9C"
RED_FILL = "FFFFC7CE"
FONT_NAME = "Arial"

MONTHLY_TARGET_PER_AGENT = 25000
WORKING_DAYS = 26
DAYS_PER_WEEK = 6  # Mon-Sat

STATUS_FORMULA = (
    'IF({pct}>=1,"\U0001F7E2 On Target",'
    'IF({pct}>=0.7,"\U0001F7E1 Watch","\U0001F534 Critical"))'
)


def style_header(cell):
    cell.font = Font(name=FONT_NAME, size=9, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_total(cell):
    cell.font = Font(name=FONT_NAME, size=9, bold=True)
    cell.fill = PatternFill("solid", fgColor=TOTAL_BLUE)
    cell.alignment = Alignment(horizontal="center", vertical="center")


def style_body(cell, bold=False, center=True):
    cell.font = Font(name=FONT_NAME, size=9, bold=bold)
    if center:
        cell.alignment = Alignment(horizontal="center", vertical="center")


def generate(agents: list[tuple[str, float]], week_ending: str, output: Path):
    n_agents = len(agents)
    per_agent_weekly_target = round(MONTHLY_TARGET_PER_AGENT / WORKING_DAYS * DAYS_PER_WEEK)
    team_weekly_target = per_agent_weekly_target * n_agents

    wb = Workbook()
    ws = wb.active
    ws.title = "Weekly Sales Report"
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:F1")
    c = ws.cell(1, 1, "TRAVNOOK TRAVEL & TOURISM — WEEKLY SALES REPORT")
    c.font = Font(name=FONT_NAME, size=14, bold=True, color=WHITE)
    c.fill = PatternFill("solid", fgColor=NAVY)
    ws.row_dimensions[1].height = 25.5

    ws.merge_cells("A2:F2")
    c = ws.cell(
        2, 1,
        f"Week Ending: {week_ending}   |   Team Lead: Bindhi   |   "
        f"Per-Agent Weekly Target: AED {per_agent_weekly_target:,.0f} (base AED {MONTHLY_TARGET_PER_AGENT:,.0f}/mo "
        f"÷ {WORKING_DAYS} working days × {DAYS_PER_WEEK})",
    )
    c.font = Font(name=FONT_NAME, size=9, color=NAVY)

    ws.merge_cells("A4:F4")
    c = ws.cell(4, 1, "▼  TEAM SUMMARY")
    c.font = Font(name=FONT_NAME, size=9, bold=True, color=NAVY)

    headers = ["TEAM TARGET", "TEAM ACHIEVED", "GAP", "% ACHIEVED", "TOP PERFORMER", "STATUS"]
    for i, h in enumerate(headers, start=1):
        style_header(ws.cell(5, i, h))
    ws.row_dimensions[5].height = 20

    first_row = 10
    last_row = first_row + n_agents - 1
    total_row = last_row + 1

    ws.cell(6, 1, team_weekly_target)
    ws.cell(6, 1).font = Font(name=FONT_NAME, size=9, color=INPUT_BLUE_FONT)
    ws.cell(6, 1).alignment = Alignment(horizontal="center")
    ws.cell(6, 2, f"=C{total_row}")
    ws.cell(6, 3, "=A6-B6")
    ws.cell(6, 4, "=IFERROR(B6/A6,0)")
    ws.cell(6, 4).number_format = "0%"
    ws.cell(6, 5, f'=INDEX(A{first_row}:A{last_row},MATCH(MAX(C{first_row}:C{last_row}),C{first_row}:C{last_row},0))')
    ws.cell(6, 6, "=" + STATUS_FORMULA.format(pct="D6"))
    for i in range(1, 7):
        style_body(ws.cell(6, i))

    ws.merge_cells("A8:F8")
    c = ws.cell(8, 1, "▼  INDIVIDUAL PERFORMANCE")
    c.font = Font(name=FONT_NAME, size=9, bold=True, color=NAVY)

    ag_headers = ["AGENT", "WEEKLY TARGET", "ACHIEVED", "GAP", "% ACHIEVED", "STATUS"]
    for i, h in enumerate(ag_headers, start=1):
        style_header(ws.cell(9, i, h))

    for i, (name, amount) in enumerate(agents):
        row = first_row + i
        style_body(ws.cell(row, 1, name), center=False)
        ws.cell(row, 2, per_agent_weekly_target)
        ws.cell(row, 3, amount)
        ws.cell(row, 4, f"=B{row}-C{row}")
        ws.cell(row, 5, f"=IFERROR(C{row}/B{row},0)")
        ws.cell(row, 5).number_format = "0%"
        ws.cell(row, 6, "=" + STATUS_FORMULA.format(pct=f"E{row}"))
        for cc in (2, 3, 4, 5, 6):
            style_body(ws.cell(row, cc))
        # conditional fill based on % achieved (computed here since openpyxl formula
        # results aren't known at write time)
        pct = amount / per_agent_weekly_target if per_agent_weekly_target else 0
        fill = GREEN_FILL if pct >= 1 else YELLOW_FILL if pct >= 0.7 else RED_FILL
        ws.cell(row, 6).fill = PatternFill("solid", fgColor=fill)

    style_total(ws.cell(total_row, 1, "TEAM TOTAL"))
    ws.cell(total_row, 2, f"=SUM(B{first_row}:B{last_row})")
    ws.cell(total_row, 3, f"=SUM(C{first_row}:C{last_row})")
    ws.cell(total_row, 4, f"=B{total_row}-C{total_row}")
    ws.cell(total_row, 5, f"=IFERROR(C{total_row}/B{total_row},0)")
    ws.cell(total_row, 5).number_format = "0%"
    ws.cell(total_row, 6, "=" + STATUS_FORMULA.format(pct=f"E{total_row}"))
    for cc in (2, 3, 4, 5, 6):
        style_total(ws.cell(total_row, cc))

    footer_row = total_row + 2
    ws.merge_cells(start_row=footer_row, start_column=1, end_row=footer_row, end_column=6)
    below_target = [name for name, amount in agents if amount / per_agent_weekly_target < 0.7]
    flag_text = (
        f"⚠️ Needs attention (below 70% of weekly target): {', '.join(below_target)}."
        if below_target else "All agents at or above 70% of weekly target."
    )
    c = ws.cell(footer_row, 1, flag_text)
    c.font = Font(name=FONT_NAME, size=9, bold=True, color="FFC00000" if below_target else GRAY)

    widths = [16, 15, 13, 13, 12, 14]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    print(f"Wrote {output}")


def parse_args():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--data", type=str, required=True, help='Comma-separated "Agent:Amount" pairs')
    p.add_argument("--week-ending", type=str, required=True)
    p.add_argument("--output", type=Path, required=True)
    return p.parse_args()


if __name__ == "__main__":
    args = parse_args()
    pairs = []
    for item in args.data.split(","):
        name, amount = item.rsplit(":", 1)
        pairs.append((name.strip(), float(amount.strip())))
    generate(pairs, args.week_ending, args.output)
