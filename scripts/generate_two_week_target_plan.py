#!/usr/bin/env python3
"""
Generate a Travnook two-week rolling target plan: monthly target, this
week's target, achieved so far, the deficit, next week's normal target,
and the recovery "aim" for next week (deficit rolled forward + next
week's own target). Mirrors the example: 25k/mo -> 5,769/wk; achieve
4,000 -> deficit 1,769 -> next week's aim = 1,769 + 5,769 = 7,538.

Usage:
    python3 generate_two_week_target_plan.py \
        --data "Harinder:12494,Ibrahim:11353,Layan (ATL):8600,Chaima:8041,Rihab:5874,Khosema:2550,Saha:2070,Alshima:0" \
        --default-target 25000 \
        --target-override "Layan (ATL):20000,Chaima:20000" \
        --this-week-label "WK2: 7-12 Sep 2026" --this-week-days 6 \
        --next-week-label "WK3: 14-19 Sep 2026" --next-week-days 6 \
        --output trackers/reports/TWO_WEEK_TARGET_PLAN_2026-09-11.xlsx
"""
import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

NAVY = "FF1F3864"
WHITE = "FFFFFFFF"
TOTAL_BLUE = "FFBDD7EE"
INPUT_BLUE_FONT = "FF0000FF"
GRAY = "FF808080"
FONT_NAME = "Arial"

WORKING_DAYS = 26

STATUS_FORMULA = (
    'IF({pct}>=1,"\U0001F7E2 On Target",'
    'IF({pct}>=0.7,"\U0001F7E1 Watch","\U0001F534 Critical"))'
)


def style_header(cell):
    cell.font = Font(name=FONT_NAME, size=9, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_total(cell):
    cell.font = Font(name=FONT_NAME, size=9, bold=True)
    cell.fill = PatternFill("solid", fgColor=TOTAL_BLUE)
    cell.alignment = Alignment(horizontal="center", vertical="center")


def style_body(cell, bold=False, center=True):
    cell.font = Font(name=FONT_NAME, size=9, bold=bold)
    if center:
        cell.alignment = Alignment(horizontal="center", vertical="center")


def generate(
    agents: list[tuple[str, float, float]],
    this_week_label: str, this_week_days: int,
    next_week_label: str, next_week_days: int,
    default_target: float, output: Path,
):
    n_agents = len(agents)
    overrides = [f"{name} AED {mt:,.0f}/mo" for name, _, mt in agents if mt != default_target]

    wb = Workbook()
    ws = wb.active
    ws.title = "Two-Week Target Plan"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A6"

    last_col = 8  # A..H

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    c = ws.cell(1, 1, "TRAVNOOK TRAVEL & TOURISM — TWO-WEEK TARGET PLAN")
    c.font = Font(name=FONT_NAME, size=14, bold=True, color=WHITE)
    c.fill = PatternFill("solid", fgColor=NAVY)
    ws.row_dimensions[1].height = 25.5

    override_note = f"   |   Target override: {', '.join(overrides)}" if overrides else ""
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    c = ws.cell(
        2, 1,
        f"This week: {this_week_label}   |   Next week: {next_week_label}   |   "
        f"Default Monthly Target: AED {default_target:,.0f}{override_note}",
    )
    c.font = Font(name=FONT_NAME, size=9, color=NAVY)

    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=last_col)
    c = ws.cell(
        3, 1,
        "Aim for next week = this week's deficit (if any) + next week's own target. A surplus this week does "
        "not reduce next week's target — it only offsets a deficit, never creates a negative one.",
    )
    c.font = Font(name=FONT_NAME, size=9, italic=True, color=GRAY)

    headers = [
        "AGENT", "MONTHLY TARGET", f"THIS WEEK TARGET\n({this_week_label})", "ACHIEVED SO FAR",
        "DEFICIT", f"NEXT WEEK TARGET\n({next_week_label})", "AIM FOR NEXT WEEK", "STATUS",
    ]
    for i, h in enumerate(headers, start=1):
        style_header(ws.cell(5, i, h))
    ws.row_dimensions[5].height = 34

    first_row = 6
    last_row = first_row + n_agents - 1
    total_row = last_row + 1

    for i, (name, achieved, monthly_target) in enumerate(agents):
        row = first_row + i
        style_body(ws.cell(row, 1, name), center=False)
        cell = ws.cell(row, 2, monthly_target)
        cell.font = Font(name=FONT_NAME, size=9, color=INPUT_BLUE_FONT)
        cell.alignment = Alignment(horizontal="center")
        ws.cell(row, 3, f"=ROUND(B{row}/{WORKING_DAYS}*{this_week_days},0)")
        ws.cell(row, 4, achieved)
        ws.cell(row, 5, f"=C{row}-D{row}")
        ws.cell(row, 6, f"=ROUND(B{row}/{WORKING_DAYS}*{next_week_days},0)")
        ws.cell(row, 7, f"=F{row}+MAX(E{row},0)")
        ws.cell(row, 8, "=" + STATUS_FORMULA.format(pct=f"IFERROR(D{row}/C{row},0)"))
        for cc in (3, 4, 5, 6, 7, 8):
            style_body(ws.cell(row, cc))

    style_total(ws.cell(total_row, 1, "TEAM TOTAL"))
    for cc in (2, 3, 4, 5, 6, 7):
        letter = get_column_letter(cc)
        ws.cell(total_row, cc, f"=SUM({letter}{first_row}:{letter}{last_row})")
        style_total(ws.cell(total_row, cc))
    ws.cell(total_row, 8, "=" + STATUS_FORMULA.format(pct=f"IFERROR(D{total_row}/C{total_row},0)"))
    style_total(ws.cell(total_row, 8))

    widths = [14, 14, 16, 14, 11, 16, 16, 13]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in range(first_row, total_row + 1):
        ws.row_dimensions[row].height = 20

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    print(f"Wrote {output}")


def parse_pairs(s: str) -> dict[str, float]:
    result = {}
    for item in s.split(","):
        item = item.strip()
        if not item:
            continue
        name, value = item.rsplit(":", 1)
        result[name.strip()] = float(value.strip())
    return result


def parse_args():
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--data", type=str, required=True)
    p.add_argument("--default-target", type=float, required=True)
    p.add_argument("--target-override", type=str, default="")
    p.add_argument("--this-week-label", type=str, required=True)
    p.add_argument("--this-week-days", type=int, default=6)
    p.add_argument("--next-week-label", type=str, required=True)
    p.add_argument("--next-week-days", type=int, default=6)
    p.add_argument("--output", type=Path, required=True)
    return p.parse_args()


if __name__ == "__main__":
    args = parse_args()
    achieved = parse_pairs(args.data)
    overrides = parse_pairs(args.target_override) if args.target_override else {}
    triples = [
        (name, amount, overrides.get(name, args.default_target))
        for name, amount in achieved.items()
    ]
    generate(
        triples, args.this_week_label, args.this_week_days,
        args.next_week_label, args.next_week_days, args.default_target, args.output,
    )
